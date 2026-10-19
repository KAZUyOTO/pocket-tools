#!/usr/bin/env python3
"""PDF -> XLSX by extracting tables (one sheet per table).

Uses pdfplumber for detection. If no tables are found, falls back to one sheet
with the page text, one line per row, so the output is never empty.

    python3 pdf_to_xlsx.py input.pdf output.xlsx
"""
import sys
import pdfplumber
from openpyxl import Workbook
from openpyxl.utils.exceptions import IllegalCharacterError


def clean(v):
    if v is None:
        return ""
    s = str(v)
    # openpyxl rejects control chars; strip them.
    return "".join(ch for ch in s if ch == "\n" or ch == "\t" or ord(ch) >= 32)


def main():
    if len(sys.argv) != 3:
        print("usage: pdf_to_xlsx.py <input.pdf> <output.xlsx>", file=sys.stderr)
        return 2
    src, dst = sys.argv[1], sys.argv[2]

    wb = Workbook()
    wb.remove(wb.active)
    found = 0

    with pdfplumber.open(src) as pdf:
        for pi, page in enumerate(pdf.pages, 1):
            for ti, table in enumerate(page.extract_tables(), 1):
                found += 1
                ws = wb.create_sheet(title=f"p{pi}_t{ti}"[:31])
                for row in table:
                    try:
                        ws.append([clean(c) for c in row])
                    except IllegalCharacterError:
                        ws.append([clean(c)[:32000] for c in row])

        if not found:
            ws = wb.create_sheet(title="text")
            for page in pdf.pages:
                for line in (page.extract_text() or "").splitlines():
                    ws.append([clean(line)])

    wb.save(dst)
    print(f"tables found: {found}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
